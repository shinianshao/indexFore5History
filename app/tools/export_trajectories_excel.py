# -*- coding: utf-8 -*-
"""将 SQLite 中的人物行迹数据库导出为格式优雅的 Excel 表格，方便用户直接查阅与校勘"""

import os
import sys
import sqlite3
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
sys.path.insert(0, ROOT)

DB_PATH = os.path.join(ROOT, "data", "index", "index.db")
OUT_XLSX = os.path.join(ROOT, "人物生平行迹数据库(61人全量10733条).xlsx")

BOOK_MAP = {
    "sj": "史記",
    "hs": "漢書",
    "hhs": "後漢書",
    "sgz": "三國志",
    "js": "晉書"
}

def export():
    print(f"正在从 {DB_PATH} 导出行迹数据库...")
    conn = sqlite3.connect(DB_PATH)
    c = conn.cursor()

    # 查出行迹数据与同期多地计算
    query = """
    SELECT 
        id, person_id, person_name, place_id, place_name, 
        time_raw, year_ad, source_book, source_chapter, 
        event_summary, coord_lng, coord_lat, sentence_uid
    FROM person_trajectories
    ORDER BY person_id, (CASE WHEN year_ad IS NULL THEN 9999 ELSE year_ad END), id
    """
    c.execute(query)
    rows = c.fetchall()

    # 预统计同期多地 (person_id, year_ad) -> set(places)
    multi_map = {}
    for r in rows:
        pid, y, pl = r[1], r[6], r[4]
        if y is not None:
            k = (pid, y)
            if k not in multi_map:
                multi_map[k] = set()
            multi_map[k].add(pl)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "人物生平行迹全量库"

    headers = [
        "序号", "历史人物", "人物ID", "地名", "地名ID", 
        "原始纪年", "公元年纪年", "同期多地", "典籍书目", "篇卷章节", 
        "事迹摘要/原典句子", "经度", "纬度", "原典句子UID"
    ]
    ws.append(headers)

    # 样式定义
    header_font = Font(name="微软雅黑", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="2B4C7E", end_color="2B4C7E", fill_type="solid")
    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")
    border_thin = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )
    fill_multi = PatternFill(start_color="FFF2E8", end_color="FFF2E8", fill_type="solid")

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = align_center

    row_num = 2
    for r in rows:
        trj_id, pid, pname, plid, plname, time_raw, year_ad, book, chapter, summary, lng, lat, uid = r
        bk_name = BOOK_MAP.get(book, book)
        
        is_multi = False
        if year_ad is not None and len(multi_map.get((pid, year_ad), set())) > 1:
            is_multi = True
        multi_tag = "同期多地" if is_multi else ""

        year_str = ""
        if year_ad is not None:
            year_str = f"前{abs(year_ad)}年" if year_ad < 0 else f"{year_ad}年"

        row_data = [
            row_num - 1,
            pname,
            pid,
            plname,
            plid,
            time_raw or "未详",
            year_str,
            multi_tag,
            bk_name,
            chapter or "",
            summary or "",
            round(lng, 4) if lng else "",
            round(lat, 4) if lat else "",
            uid
        ]
        ws.append(row_data)

        # 单元格格式化
        for col_idx in range(1, len(row_data) + 1):
            cell = ws.cell(row=row_num, column=col_idx)
            cell.border = border_thin
            if col_idx in [1, 2, 4, 6, 7, 8, 9, 12, 13]:
                cell.alignment = align_center
            else:
                cell.alignment = align_left
            if is_multi and col_idx == 8:
                cell.fill = fill_multi
                cell.font = Font(name="微软雅黑", size=10, bold=True, color="C00000")

        row_num += 1

    # 开启自动筛选
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{row_num - 1}"

    # 冻结首行
    ws.freeze_panes = "A2"

    # 设置列宽
    col_widths = {
        1: 8,   # 序号
        2: 12,  # 人物
        3: 14,  # 人物ID
        4: 12,  # 地名
        5: 14,  # 地名ID
        6: 18,  # 原始纪年
        7: 12,  # 公元年
        8: 12,  # 同期多地
        9: 10,  # 典籍
        10: 22, # 章节
        11: 50, # 事迹摘要
        12: 10, # 经度
        13: 10, # 纬度
        14: 16  # UID
    }
    for c_idx, w in col_widths.items():
        ws.column_dimensions[get_column_letter(c_idx)].width = w

    # 使用 safe_save_workbook 原子写入
    from pipeline.common import safe_save_workbook
    safe_save_workbook(wb, OUT_XLSX)
    conn.close()

    size_mb = os.path.getsize(OUT_XLSX) / (1024 * 1024)
    print(f"✓ 成功导出 Excel: {OUT_XLSX} (大小: {size_mb:.2f} MB, 总行数: {row_num - 1})")

if __name__ == "__main__":
    export()
