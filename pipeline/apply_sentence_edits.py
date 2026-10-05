# -*- coding: utf-8 -*-
"""P4-1 句级编辑：拆句 / 并句 / 弃用，改的是语料层（data/corpus/*.json）。

为什么改语料而不是改索引
------------------------
句子边界是 build.py 切出来的，annotate 按切分结果标注。若在索引层改边界，
下次 annotate 一跑就又被切回原样。所以编辑必须落在**语料**，让它成为新的事实。

uid 的三条继承规则（红线 2）
---------------------------
| 操作 | 处理 |
|---|---|
| 拆一句为二 | 前半**继承原 uid**，后半生成新 uid（md5 由原 uid + 断点派生） |
| 合并二句 | 保留第一句 uid，第二句标 `status=merged`（**不物理删**） |
| 弃用 | 标 `status=dead`（**不物理删**） |

被合并/弃用的句子仍留在语料里，只是 annotate 跳过它们（`status != active`），
这样历史 uid 永远查得到出处。

幂等
----
应用过的编辑会把 edit_id 记进句子的 `edits` 列表，重复跑会跳过——
否则第二次跑会把同一句再拆一次。

用法
----
    python pipeline/apply_sentence_edits.py init
    python pipeline/apply_sentence_edits.py show --uid <uid>      # 看这句原文
    python pipeline/apply_sentence_edits.py add --uid <uid> --action split --at 12
    python pipeline/apply_sentence_edits.py add --uid <uid> --action merge
    python pipeline/apply_sentence_edits.py add --uid <uid> --action dead
    python pipeline/apply_sentence_edits.py list
    python pipeline/apply_sentence_edits.py revoke --uid <uid>
    python pipeline/apply_sentence_edits.py apply [--dry-run]
"""
from __future__ import annotations

import argparse
import glob
import hashlib
import json
import os
import shutil
import sqlite3
import sys
from datetime import datetime

sys.stdout.reconfigure(encoding="utf-8")

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, HERE)
from common import CORPUS, stable_uid, safe_save_workbook   # noqa: E402

DB_PATH = os.path.join(ROOT, "data", "index", "index.db")
WORKBOOK = os.path.join(ROOT, "workbook", "sentence-edits.xlsx")
ORIG_DIR = os.path.join(ROOT, "data", "corpus-orig")   # 原始副本，供重放/还原

HEADERS = ["uid", "动作(action)", "位置(at)", "状态(status)",
           "备注(note)", "录入时间(createdAt)"]
ACTIONS = ("split", "merge", "dead")


def _openpyxl():
    try:
        import openpyxl
        return openpyxl
    except ImportError:
        raise SystemExit("缺少 openpyxl：`python -m pip install openpyxl`")


def now() -> str:
    return datetime.now().strftime("%Y-%m-%d %H:%M:%S")


# ---------------------------------------------------------------- 表

def init_workbook() -> None:
    opx = _openpyxl()
    if os.path.exists(WORKBOOK):
        return
    wb = opx.Workbook()
    ws = wb.active
    ws.title = "edits"
    ws.append(HEADERS)
    os.makedirs(os.path.dirname(WORKBOOK), exist_ok=True)
    safe_save_workbook(wb, WORKBOOK)
    print("已建空表：{}".format(WORKBOOK))


def _read_rows():
    """读回生效中的编辑行（只读）。"""
    if not os.path.exists(WORKBOOK):
        return []
    opx = _openpyxl()
    wb = opx.load_workbook(WORKBOOK, read_only=True, data_only=True)
    ws = wb["edits"] if "edits" in wb.sheetnames else wb.active
    out = []
    for i, r in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if not r or not r[0]:
            continue
        d = dict(zip(HEADERS, list(r) + [None] * (len(HEADERS) - len(r))))
        if str(d.get("状态(status)") or "active").strip().lower() != "active":
            continue
        d["_row"] = i          # 行号即 edit_id，幂等靠它
        out.append(d)
    wb.close()
    return out


def _append_row(row: dict) -> None:
    opx = _openpyxl()
    init_workbook()
    wb = opx.load_workbook(WORKBOOK)
    ws = wb["edits"] if "edits" in wb.sheetnames else wb.active
    ws.append([row.get(h) for h in HEADERS])
    try:
        safe_save_workbook(wb, WORKBOOK)
        print("已写入：{}".format(WORKBOOK))
    except (PermissionError, OSError):
        alt = WORKBOOK.replace(".xlsx", ".new.xlsx")
        try:
            safe_save_workbook(wb, alt)
            print("⚠️ 原表被占用，已改写：{}".format(alt))
        except Exception as e:
            print("⚠️ 写入备用表失败：{}".format(e))


def _set_status(uid: str, status: str) -> None:
    """作废某 uid 的所有编辑（状态改 dead，不删行）。"""
    opx = _openpyxl()
    if not os.path.exists(WORKBOOK):
        raise SystemExit("还没有 {}".format(WORKBOOK))
    wb = opx.load_workbook(WORKBOOK)
    ws = wb["edits"] if "edits" in wb.sheetnames else wb.active
    col = HEADERS.index("状态(status)") + 1
    n = 0
    for r in ws.iter_rows(min_row=2):
        if str(r[0].value or "").strip() != uid:
            continue
        if str(r[col - 1].value or "").strip() == "dead":
            continue
        ws.cell(row=r[0].row, column=col, value="dead")
        n += 1
    if not n:
        print("没有匹配的生效编辑：{}".format(uid))
        return
    try:
        safe_save_workbook(wb, WORKBOOK)
        print("已作废 {} 条（{}）".format(n, uid))
    except (PermissionError, OSError):
        alt = WORKBOOK.replace(".xlsx", ".new.xlsx")
        try:
            safe_save_workbook(wb, alt)
            print("⚠️ 原表被占用，已改写 .new.xlsx")
        except Exception as e:
            print("⚠️ 写入备用表失败：{}".format(e))



# ---------------------------------------------------------------- 定位

def uid_index() -> dict:
    """uid → 语料文件路径。564 篇全扫一遍（几秒），换来不用维护索引文件。"""
    out = {}
    for path in sorted(glob.glob(os.path.join(CORPUS, "*.json"))):
        with open(path, encoding="utf-8") as f:
            d = json.load(f)
        for p in d.get("paragraphs") or []:
            for s in p.get("sentences") or []:
                if s.get("uid"):
                    out[s["uid"]] = path
    return out


def find_sentence(doc, uid):
    """在语料文档里找句子，返回 (paragraph, index, sentence)。"""
    for p in doc.get("paragraphs") or []:
        for i, s in enumerate(p.get("sentences") or []):
            if s.get("uid") == uid:
                return p, i, s
    return None, -1, None


def new_uid_for_split(uid: str, at: int) -> str:
    """拆出来的后半句的新 uid：由原 uid 派生，保证可复现、不撞车。"""
    return hashlib.md5("{}|split|{}".format(uid, at).encode("utf-8")).hexdigest()[:12]


# ---------------------------------------------------------------- 应用

def apply_one(doc, edit) -> str:
    """对一份语料文档套用一条编辑。返回 'ok' / 'skip' / 'miss'。"""
    uid = str(edit["uid"]).strip()
    action = str(edit["动作(action)"] or "").strip()
    eid = edit["_row"]
    para, idx, sent = find_sentence(doc, uid)
    if sent is None:
        return "miss"
    edits = sent.get("edits") or []
    if eid in edits:
        return "skip"

    if action == "dead":
        sent["status"] = "dead"
    elif action == "merge":
        nxt = None
        for s in (para.get("sentences") or [])[idx + 1:]:
            if s.get("status", "active") == "active":
                nxt = s
                break
        if nxt is None:
            return "miss"
        sent["text"] = sent["text"] + nxt["text"]
        nxt["text"] = ""                     # 内容并走，本体留着可追溯
        nxt["status"] = "merged"
        nxt["merged_into"] = uid
        nxt["edits"] = (nxt.get("edits") or []) + [eid]
    elif action == "split":
        at = int(edit["位置(at)"] or 0)
        text = sent["text"]
        if not (0 < at < len(text)):
            return "miss"
        tail = {
            "id": "{}-{}".format(sent.get("id"), at),
            "seq": (sent.get("seq") or 0) + 0.5,   # 浮点插在中间，不打乱后续整数序号
            "uid": new_uid_for_split(uid, at),
            "text": text[at:],
            "edits": [eid],
        }
        sent["text"] = text[:at]             # 前半**继承原 uid**
        sents = para["sentences"]
        sents.insert(idx + 1, tail)
    else:
        return "miss"

    sent["edits"] = edits + [eid]
    return "ok"


def _backup_path(path: str) -> str:
    """原始副本放在 data/corpus-orig/ 下（与语料同名）。"""
    return os.path.join(ORIG_DIR, os.path.basename(path))


def cmd_apply(args) -> int:
    """从**原始副本重放**全部生效编辑。

    为什么不就地改：就地改是破坏性的，一旦 revoke 就回不去（拆开的句合不回来）。
    改成「每次都从原始副本重放」，撤销一条后重跑，那篇自然回到原样——
    撤销能力就是这么来的，不需要额外的回滚代码。
    """
    rows = _read_rows()
    os.makedirs(ORIG_DIR, exist_ok=True)
    index = uid_index()

    # 目标篇 = 有编辑的篇 ∪ 曾有编辑的篇（后者要**还原**）
    targets = {}
    for e in rows:
        path = index.get(str(e["uid"]).strip())
        if path:
            targets.setdefault(path, []).append(e)
    for name in os.listdir(ORIG_DIR):
        p = os.path.join(CORPUS, name)
        if os.path.exists(p):
            targets.setdefault(p, [])

    if not targets:
        return 0

    stat = {"ok": 0, "skip": 0, "miss": 0}
    missed = []
    for path, edits in targets.items():
        bpath = _backup_path(path)
        if not os.path.exists(bpath):
            shutil.copyfile(path, bpath)     # 首次编辑：先把原始存下来
        with open(bpath, encoding="utf-8") as f:
            doc = json.load(f)
        for e in edits:
            r = apply_one(doc, e)
            stat[r] += 1
            if r == "miss":
                missed.append("{} [{}]".format(e["uid"], e["动作(action)"]))
        if not args.dry_run:
            with open(path, "w", encoding="utf-8") as f:
                json.dump(doc, f, ensure_ascii=False, indent=1)

    print("句级编辑套用：生效 {} / 已应用跳过 {} / 未匹配 {}{}".format(
        stat["ok"], stat["skip"], stat["miss"],
        "　{}".format(" ".join(missed[:8])) if missed else ""))
    print("  重放 {} 篇语料（原始副本 {} 份）{}".format(
        len(targets), len(os.listdir(ORIG_DIR)),
        "　[dry-run 未写入]" if args.dry_run else ""))
    return stat["ok"]


def cmd_show(args) -> None:
    if not os.path.exists(DB_PATH):
        raise SystemExit("索引库不存在，先跑 rebuild.py")
    conn = sqlite3.connect(DB_PATH)
    row = conn.execute(
        "SELECT s.text, s.chapter_id, s.para_seq, s.seq FROM sentences s "
        "WHERE s.uid=?", (args.uid,)).fetchone()
    conn.close()
    if not row:
        raise SystemExit("库里没有这个 uid：{}".format(args.uid))
    text, cid, para, seq = row
    print("{} · {}（段 {} 句 {}）\n  {}".format(cid, args.uid, para, seq, text))
    print("\n共 {} 字。拆句示例：--action split --at {}".format(len(text), len(text) // 2))


def cmd_add(args) -> None:
    if args.action == "split" and not args.at:
        raise SystemExit("split 必须给 --at（在第几个字后断开，1-based）")
    _append_row({
        "uid": args.uid,
        "动作(action)": args.action,
        "位置(at)": args.at or "",
        "状态(status)": "active",
        "备注(note)": args.note or "",
        "录入时间(createdAt)": now(),
    })


def cmd_list(args) -> None:
    rows = _read_rows()
    if not rows:
        print("（还没有句级编辑）")
        return
    print("=== 句级编辑（{} 条生效中）===".format(len(rows)))
    for r in rows:
        print("  #{} {} [{}] at={}　{}".format(
            r["_row"], r["uid"], r["动作(action)"], r["位置(at)"] or "-",
            r["备注(note)"] or ""))


def main() -> int:
    ap = argparse.ArgumentParser(description="BOOKINDEX 句级编辑（拆/并/弃用）")
    sub = ap.add_subparsers(dest="cmd", required=True)

    sub.add_parser("init", help="建空表").set_defaults(
        func=lambda a: init_workbook())

    p = sub.add_parser("show", help="看某句原文（含字数，方便定 --at）")
    p.add_argument("--uid", required=True)
    p.set_defaults(func=cmd_show)

    p = sub.add_parser("add", help="写一条编辑")
    p.add_argument("--uid", required=True)
    p.add_argument("--action", required=True, choices=ACTIONS)
    p.add_argument("--at", type=int, default=0, help="split 用：在第几个字后断开")
    p.add_argument("--note", default="")
    p.set_defaults(func=cmd_add)

    sub.add_parser("list", help="列已有编辑").set_defaults(func=cmd_list)

    p = sub.add_parser("revoke", help="作废某 uid 的编辑（状态改 dead）")
    p.add_argument("--uid", required=True)
    p.set_defaults(func=lambda a: _set_status(a.uid, "dead"))

    p = sub.add_parser("apply", help="套用到语料")
    p.add_argument("--dry-run", action="store_true")
    p.set_defaults(func=cmd_apply)

    args = ap.parse_args()
    args.func(args)
    return 0


if __name__ == "__main__":
    sys.exit(main())
