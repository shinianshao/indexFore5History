# pipeline/_fix_persons_accuracy.py
# -*- coding: utf-8 -*-
"""修复 persons.xlsx 中周平王别名、朝代元数据错误及高危短别名"""
import os
import sys
import openpyxl

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import common

def main():
    root = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    xlsx_path = os.path.join(root, "workbook", "persons.xlsx")

    wb = openpyxl.load_workbook(xlsx_path)
    ws = wb["人物"]

    header = [c for c in next(ws.iter_rows(values_only=True))]
    id_col = header.index("id")
    dyn_col = header.index("朝代")
    alias_col = header.index("别名(竖线分隔)")
    in_sj_col = header.index("in_sj")
    in_hs_col = header.index("in_hs")
    in_hhs_col = header.index("in_hhs")
    in_sgz_col = header.index("in_sgz")
    in_js_col = header.index("in_js")

    modifications = {
        # 1. 周平王：移除裸“平王”与裸“宜臼”，限定四书（排除sgz）
        "p_zhoupingwang": {
            "dynasty": "東周",
            "aliases": "周平王|姬宜臼",
            "in_sgz": 0,
        },
        # 2. 潘岳：修复朝代为西晋
        "p_panyue": {
            "dynasty": "西晉",
            "aliases": "潘岳|潘安|潘安仁",
            "in_js": 1,
        },
        # 3. 挚虞：修复朝代为西晋
        "p_zhiyu": {
            "dynasty": "西晉",
            "aliases": "摯虞|挚虞|仲洽|摯仲洽",
            "in_js": 1,
        },
        # 4. 嵇康：修复朝代为三国
        "p_jikang": {
            "dynasty": "三國",
            "aliases": "嵇康|嵇叔夜|叔夜",
            "in_sgz": 1,
            "in_js": 1,
        },
        # 5. 郝昭：移除“伯道”，杜绝汉书五行志“齐桓行伯道”11处伪命中
        "p_haozhao": {
            "aliases": "郝昭",
            "in_sgz": 1,
        },
        # 6. 夏方：移除“文正”，杜绝汉书“考文正理”伪命中
        "p_xiafang": {
            "aliases": "夏方",
            "in_js": 1,
        },
        # 7. 毕卓：移除“茂世”，杜绝汉书“丰茂世之规”伪命中
        "p_bizhuo": {
            "aliases": "畢卓|毕卓",
            "in_js": 1,
        },
        # 8. 蔡豹：移除“士宣”，杜绝汉书“公士宣”伪命中
        "p_caibao": {
            "aliases": "蔡豹",
            "in_js": 1,
        },
        # 9. 何曾：去重别名，限定三国两晋
        "p_hezeng": {
            "dynasty": "西晉",
            "aliases": "何曾|穎考|何穎考",
            "in_sj": 0,
            "in_hs": 0,
            "in_hhs": 0,
            "in_sgz": 1,
            "in_js": 1,
        },
        # 10. 顾众：限定晋书
        "p_guzhong": {
            "aliases": "顧衆|顾众|顧長始",
            "in_sj": 0,
            "in_hs": 0,
            "in_hhs": 0,
            "in_sgz": 0,
            "in_js": 1,
        },
        # 11. 孙辅：限定三国志
        "p_sunfu": {
            "in_sj": 0,
            "in_hs": 0,
            "in_hhs": 0,
            "in_sgz": 1,
            "in_js": 0,
        },
        # 12. 孙和：限定三国志
        "p_sunhe": {
            "in_sj": 0,
            "in_hs": 0,
            "in_hhs": 0,
            "in_sgz": 1,
            "in_js": 0,
        },
        # 13. 孙奋：限定三国志
        "p_sunfen": {
            "in_sj": 0,
            "in_hs": 0,
            "in_hhs": 0,
            "in_sgz": 1,
            "in_js": 0,
        },
    }

    updated = 0
    for row_idx in range(2, ws.max_row + 1):
        pid = ws.cell(row=row_idx, column=id_col + 1).value
        if pid in modifications:
            m = modifications[pid]
            if "dynasty" in m:
                ws.cell(row=row_idx, column=dyn_col + 1).value = m["dynasty"]
            if "aliases" in m:
                ws.cell(row=row_idx, column=alias_col + 1).value = m["aliases"]
            if "in_sj" in m:
                ws.cell(row=row_idx, column=in_sj_col + 1).value = m["in_sj"]
            if "in_hs" in m:
                ws.cell(row=row_idx, column=in_hs_col + 1).value = m["in_hs"]
            if "in_hhs" in m:
                ws.cell(row=row_idx, column=in_hhs_col + 1).value = m["in_hhs"]
            if "in_sgz" in m:
                ws.cell(row=row_idx, column=in_sgz_col + 1).value = m["in_sgz"]
            if "in_js" in m:
                ws.cell(row=row_idx, column=in_js_col + 1).value = m["in_js"]
            print(f"✓ 已更新 {pid}: {m}")
            updated += 1

    print(f"\n总计更新 {updated} 位人物的权威源属性。")
    common.safe_save_workbook(wb, xlsx_path)
    print(f"✓ 权威源 {xlsx_path} 已通过 safe_save_workbook 安全保存！")

if __name__ == "__main__":
    main()
