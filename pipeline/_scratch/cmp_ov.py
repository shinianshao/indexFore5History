"""临时：逐行比对 workbook/overrides.xlsx 与 HEAD 版本，确认是openpyxl 重存的字节噪声。"""
import io
import subprocess
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[2]
rel = sys.argv[1]
p = ROOT / rel

blob = subprocess.run(["git", "show", "HEAD:" + rel], cwd=str(ROOT),
                      capture_output=True).stdout
tmp = Path("C:/Users/dell/AppData/Local/Temp/bi_ov_head.xlsx")
tmp.write_bytes(blob)

a = openpyxl.load_workbook(str(tmp), data_only=True)
b = openpyxl.load_workbook(str(p), data_only=True)

print("sheet 名:", a.sheetnames, "vs", b.sheetnames)
if a.sheetnames != b.sheetnames:
    print("✘ sheet 名不同")
    sys.exit(1)

diff = 0
for name in a.sheetnames:
    sa, sb = a[name], b[name]
    if (sa.max_row, sa.max_column) != (sb.max_row, sb.max_column):
        print("✘ {} 尺寸不同 {} vs {}".format(name, (sa.max_row, sa.max_column),
                                       (sb.max_row, sb.max_column)))
        diff += 1
        continue
    for ra, rb in zip(sa.iter_rows(values_only=True),
                      sb.iter_rows(values_only=True)):
        if ra != rb:
            diff += 1
            if diff <= 5:
                print("  差异行 HEAD:", ra)
                print("        工作区:", rb)
print("总差异行数:", diff)
if diff == 0:
    print("==> 内容逐行相同，是 openpyxl 重存的字节噪声，可安全 git checkout --")
else:
    print("==> 内容真的变了，不要还原")
